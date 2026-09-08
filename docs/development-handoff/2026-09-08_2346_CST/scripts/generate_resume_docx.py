from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path('/home/ubuntu/mindscribe-live-repo')
OUT = ROOT / 'app/documents/resumes'
OUT.mkdir(parents=True, exist_ok=True)
ASSETS = Path('/home/ubuntu/liumoxin-resume-typst/assets')
PHOTO = ASSETS / 'profile.png'
LOGO = ASSETS / 'mindscribe-brand.png'
PORTFOLIO = 'https://mindscribe-app.pages.dev/'

NAVY = '101D35'
BLUE = '0057D9'
MUTED = '5B677A'
PALE = 'EEF4FF'
RULE = 'CCD7E8'
WHITE = 'FFFFFF'
BODY = '1B2330'


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_layout(table) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')


def set_paragraph_border(paragraph, color=BLUE, size='10', space='2') -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr')
        p_pr.append(borders)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)
    borders.append(bottom)


def set_run_font(run, size: float, bold=False, color=BODY, latin='Arial', cjk='Microsoft YaHei') -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement('w:rFonts')
        r_pr.append(r_fonts)
    r_fonts.set(qn('w:ascii'), latin)
    r_fonts.set(qn('w:hAnsi'), latin)
    r_fonts.set(qn('w:eastAsia'), cjk)


def tune_paragraph(paragraph, before=0, after=0, line=1.0, keep=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True


def clear_cell(cell) -> None:
    cell.text = ''
    p = cell.paragraphs[0]
    tune_paragraph(p)


def add_hyperlink(paragraph, text: str, url: str, size=8.0, color=BLUE, bold=False) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    r_fonts = OxmlElement('w:rFonts')
    r_fonts.set(qn('w:ascii'), 'Arial')
    r_fonts.set(qn('w:hAnsi'), 'Arial')
    r_fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    r_pr.append(r_fonts)
    color_node = OxmlElement('w:color')
    color_node.set(qn('w:val'), color)
    r_pr.append(color_node)
    size_node = OxmlElement('w:sz')
    size_node.set(qn('w:val'), str(int(size * 2)))
    r_pr.append(size_node)
    if bold:
        r_pr.append(OxmlElement('w:b'))
    new_run.append(r_pr)
    text_node = OxmlElement('w:t')
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def section_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    tune_paragraph(p, before=5.2, after=3.0, line=1.0, keep=True)
    set_paragraph_border(p)
    set_run_font(p.add_run(title), 10.2, True, NAVY)


def role_heading(doc: Document, company: str, role: str, period: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_repeat_table_layout(table)
    table.columns[0].width = Cm(13.8)
    table.columns[1].width = Cm(4.2)
    left, right = table.rows[0].cells
    clear_cell(left)
    clear_cell(right)
    p = left.paragraphs[0]
    set_run_font(p.add_run(company), 9.1, True, NAVY)
    set_run_font(p.add_run('  ' + role), 8.3, True, BLUE)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(rp.add_run(period), 7.7, True, MUTED)


def body_text(doc: Document, text: str, size=8.0, color=BODY, before=0, after=1.2, bold=False) -> None:
    p = doc.add_paragraph()
    tune_paragraph(p, before=before, after=after, line=1.03)
    set_run_font(p.add_run(text), size, bold, color)


def bullet_list(doc: Document, items: Iterable[str], size=7.55) -> None:
    for item in items:
        p = doc.add_paragraph()
        tune_paragraph(p, after=0.7, line=1.01)
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.23)
        set_run_font(p.add_run('•  '), size, True, BLUE)
        set_run_font(p.add_run(item), size, False, BODY)


def metric_row(doc: Document, metrics: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(metrics))
    set_repeat_table_layout(table)
    for i, (value, label) in enumerate(metrics):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, PALE)
        set_cell_margins(cell, 65, 75, 65, 75)
        clear_cell(cell)
        p = cell.paragraphs[0]
        set_run_font(p.add_run(value), 11.0, True, BLUE)
        p.add_run('\n')
        set_run_font(p.add_run(label), 6.9, True, MUTED)


def product_matrix(doc: Document, items: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(items))
    set_repeat_table_layout(table)
    for i, (title, detail) in enumerate(items):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, PALE)
        set_cell_margins(cell, 110, 85, 110, 85)
        clear_cell(cell)
        p = cell.paragraphs[0]
        set_run_font(p.add_run(title), 8.3, True, NAVY)
        p.add_run('\n')
        set_run_font(p.add_run(detail), 6.8, False, MUTED)


def target_footer(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    tune_paragraph(p, before=4.0, after=0, line=1.0)
    set_paragraph_border(p, color=RULE, size='6', space='2')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), 6.9, False, MUTED)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def enlarge_document_typography(doc: Document, factor: float = 1.08) -> None:
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(iter_table_paragraphs(table))
    for paragraph in paragraphs:
        fmt = paragraph.paragraph_format
        if isinstance(fmt.line_spacing, float):
            fmt.line_spacing = max(1.05, fmt.line_spacing * 1.03)
        if fmt.space_after is not None and fmt.space_after.pt > 0:
            fmt.space_after = Pt(fmt.space_after.pt * 1.08)
        for run in paragraph.runs:
            if run.font.size is not None:
                run.font.size = Pt(run.font.size.pt * factor)


def dark_bar(doc: Document, lead: str, detail: str, link=False) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_repeat_table_layout(table)
    table.columns[0].width = Cm(11.5)
    table.columns[1].width = Cm(7.0)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 70, 95, 70, 95)
    left, right = table.rows[0].cells
    clear_cell(left)
    clear_cell(right)
    set_run_font(left.paragraphs[0].add_run(lead), 7.4, True, WHITE)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if link:
        add_hyperlink(right.paragraphs[0], detail, PORTFOLIO, 7.4, '8BB0FF', True)
    else:
        set_run_font(right.paragraphs[0].add_run(detail), 7.4, False, 'DCE6F7')


def add_header(doc: Document, *, name: str, title: str, subtitle: str | None, location: str, email: str, portfolio_label: str, language_line: str) -> None:
    table = doc.add_table(rows=1, cols=3)
    set_repeat_table_layout(table)
    widths = [2.5, 9.4, 6.5]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    photo_cell, name_cell, contact_cell = table.rows[0].cells
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 20, 0, 20)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        clear_cell(cell)
    photo_cell.paragraphs[0].add_run().add_picture(str(PHOTO), width=Cm(2.05), height=Cm(2.65))
    np = name_cell.paragraphs[0]
    add_hyperlink(np, name, PORTFOLIO, 20.5, NAVY, True)
    np.add_run('\n')
    set_run_font(np.add_run(title), 11.2, True, BLUE)
    if subtitle:
        np.add_run('\n')
        set_run_font(np.add_run(subtitle), 7.4, True, MUTED)
    cp = contact_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cp.add_run().add_picture(str(LOGO), width=Cm(0.95))
    cp.add_run('\n')
    set_run_font(cp.add_run('MindScribe'), 6.9, True, NAVY)
    cp.add_run('\n')
    set_run_font(cp.add_run(location), 7.5, False, NAVY)
    cp.add_run('\n')
    add_hyperlink(cp, email, 'mailto:' + email, 7.3)
    cp.add_run('\n')
    add_hyperlink(cp, portfolio_label, PORTFOLIO, 7.0)
    cp.add_run('\n')
    set_run_font(cp.add_run(language_line), 7.0, False, MUTED)


def add_page2_header(doc: Document, name: str, title: str, contact: str) -> None:
    table = doc.add_table(rows=1, cols=3)
    set_repeat_table_layout(table)
    table.columns[0].width = Cm(1.0)
    table.columns[1].width = Cm(11.4)
    table.columns[2].width = Cm(6.0)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 15, 0, 15)
        clear_cell(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(LOGO), width=Cm(0.62))
    p = table.cell(0, 1).paragraphs[0]
    set_run_font(p.add_run(name), 11.2, True, NAVY)
    set_run_font(p.add_run('  ' + title), 7.7, True, BLUE)
    p = table.cell(0, 2).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run(contact), 6.9, False, MUTED)


def add_profile_box(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE)
    set_cell_margins(cell, 75, 100, 75, 100)
    clear_cell(cell)
    p = cell.paragraphs[0]
    set_run_font(p.add_run(label + '  '), 8.0, True, NAVY)
    set_run_font(p.add_run(text), 7.55, False, BODY)


def add_caps(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    tune_paragraph(p, after=0, line=1.0)
    set_run_font(p.add_run(text), 7.0, True, BLUE)


def add_case(doc: Document, number_title: str, text: str, meta: str | None = None) -> None:
    p = doc.add_paragraph()
    tune_paragraph(p, before=1.8, after=0.5, line=1.0, keep=True)
    set_run_font(p.add_run(number_title), 8.25, True, NAVY)
    if meta:
        p = doc.add_paragraph()
        tune_paragraph(p, after=0.5, line=1.0, keep=True)
        set_run_font(p.add_run(meta), 7.0, False, MUTED)
    body_text(doc, text, 7.45, BODY, after=1.0)


def build_cn() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.0)
    sec.bottom_margin = Cm(0.9)
    sec.left_margin = Cm(1.25)
    sec.right_margin = Cm(1.25)
    sec.header_distance = Cm(0)
    sec.footer_distance = Cm(0)
    styles = doc.styles
    styles['Normal'].font.name = 'Microsoft YaHei'
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    styles['Normal'].font.size = Pt(8)

    add_header(doc, name='刘莫昕', title='智能硬件项目经理', subtitle='SMART HARDWARE PROJECT MANAGER', location='杭州 · 159 9446 1673', email='scretliasion@163.com', portfolio_label='作品集：mindscribe-app.pages.dev', language_line='英语可作为工作语言')
    body_text(doc, '', after=1)
    add_profile_box(doc, '职业概述', '具备门禁机、家用门锁、玻璃门锁及房车锁多产品线全生命周期管理经验，覆盖产品定义、研发跟进、测试验证、认证合规、供应商协同、试产量产与售后闭环。硕士阶段接受全英文教学，英语可作为工作语言；具备独立对接尼泊尔等海外客户的需求澄清与项目沟通经验。能够在软硬件、结构、固件、云平台与供应链之间建立共同语言，以里程碑、风险清单和问题闭环推动稳定交付。')
    section_heading(doc, '核心能力')
    add_caps(doc, '全生命周期管理  ·  研发与供应商协同  ·  软硬件联调  ·  测试验证  ·  认证合规  ·  试产与量产  ·  客诉闭环  ·  英语项目沟通')
    section_heading(doc, '工作经历')
    role_heading(doc, '得力集团', '智能硬件项目经理', '2026.03 — 至今')
    body_text(doc, '负责门禁机、家用门锁、玻璃门锁与房车锁产品线，统筹自研整机与 OEM 项目的研发、测试、标准化、交付和售后。', 7.55, MUTED, after=1.0)
    bullet_list(doc, [
        '并行推进 2.4 英寸、5 英寸、7 英寸门禁机及多款 OEM 机型；完成 AL912C、AC213、AC205C 三款 2.4 英寸门禁机研发、测试与试产，按计划于 2026 年 8 月上市。',
        '统筹 5 英寸防水门禁机 AC516/AL962C 的型式测试、硬件可靠性验证、首轮功能测试及包材设计；针对掌静脉识别与物料交期问题组织复测、分析与风险复盘。',
        '负责 AC22 人脸家用门锁新国标适配，对比检测机构方案并推进认证；同步完成 DL-AC27 等存量机型标准差异梳理、参数与资料迭代。',
        '全流程推进房车锁项目，覆盖专利风险排查、供应商寻源、价格核定、试产验证与量产驻场；推动供应商解决低温锁舌失效，实现样机在 -30℃ 场景下正常动作。',
        '接手门禁与锁具售后台账后，完成 10 条历史问题中的 5 条闭环；组织继电器虚焊、指纹识别、结构异响、门铃无响应等问题的根因分析与验证结案。',
        '针对 AC512 门铃无响应问题拆解电路原理与典型错误接线场景，输出专项失效分析材料，为售后判定和现场整改建立依据。',
        '推动全系门禁机参数、操作说明和技术规范对齐，补充操作视频二维码；协同售后团队明确岗位边界、问题流转与结案标准。',
        '独立对接尼泊尔、巴基斯坦及其他海外区域客户，完成定制需求澄清、技术方案同步与项目进度沟通。',
    ])
    dark_bar(doc, '交付方法', '里程碑拆解 → 风险清单 → 责任人 → 验证标准 → 闭环复盘')
    role_heading(doc, '广西诺特广电有限公司', '热融合设备实习项目经理', '2024.06 — 2024.09')
    bullet_list(doc, ['协调设计团队与工厂资源，参与产品开发全周期，推动 3 款新品由设计阶段进入试产。'], 7.45)
    product_matrix(doc, [
        ('门禁终端', '2.4 / 5 / 7 英寸与 OEM'),
        ('家用门锁', '整机交付与合规适配'),
        ('玻璃门锁', '识别、交互与项目协同'),
        ('房车锁', '结构、试产与量产跟进'),
    ])
    dark_bar(doc, '项目交付链路', '产品定义 → 研发协同 → 测试认证 → 供应链 → 试产量产 → 售后闭环')
    target_footer(doc, '求职方向：智能硬件项目经理｜产品项目经理｜智能门禁 / IoT / 机器人与自动化相关方向')

    doc.add_page_break()
    add_page2_header(doc, '刘莫昕', '智能硬件项目经理', '159 9446 1673 · scretliasion@163.com')
    metric_row(doc, [('3 款', '2.4 英寸门禁机如期上市'), ('5 / 10', '历史客诉完成闭环'), ('-30℃', '房车锁低温验证通过')])
    section_heading(doc, '代表项目与业务成果')
    add_case(doc, '01｜门禁机多型号上市与 OEM 交付', '从项目策划、协议与功能测试、可靠性验证到试产和入库，拆解关键节点并组织研发、品质、采购、供应商及生产协同；既管理自研整机，也处理 OEM 固件兼容、资料对齐与交付风险。', 'AL912C / AC213 / AC205C / AC516 / AL962C / AC611 / AC711 / AC612 / AC615')
    add_case(doc, '02｜智能门锁新国标切换', '围绕 GB 21556.2-2025 切换窗口，推进 AC22 新品认证与 AC10、AC21、AC310、DL-AC27 等在售产品升级；完成检测方案比较、测试不符合项整改、标准差异核对与资料更新。')
    add_case(doc, '03｜品质问题与客诉闭环', '建立问题台账，按“现象复现—影响评估—软硬件定位—责任与方案—验证结案”推进闭环；覆盖主板虚焊、固件异常、识别效果、结构缺陷、电源适配和现场接线等典型问题。')
    section_heading(doc, '教育背景')
    role_heading(doc, '浙江工商大学萨塞克斯人工智能学院', '机器人与自动化系统（全英文授课）｜硕士', '')
    body_text(doc, '机器学习与深度学习、计算机视觉、运动规划与控制、嵌入式人工智能、三维建模与力学仿真', 7.35, MUTED)
    role_heading(doc, '浙江工商大学', '电子信息专业｜本科', '')
    body_text(doc, '电路分析、数字/模拟电子技术、单片机综合、信号与系统、EDA 技术', 7.35, MUTED)
    section_heading(doc, '工程与研究项目')
    bullet_list(doc, [
        '高铁 5G 波束与资源协同：构建四场景通信仿真平台，以 PPO 与 Q-learning 协同处理连续波束/功率和离散切换；报告仿真中直线轨道切换成功率由 92.8% 提升至 97.4%。',
        '无人机蜂窝覆盖盲点搜索：设计航迹规划、信道模型与信号可视化三模块系统；内螺旋航迹较往复式航迹模拟用时减少约 28.92%—33.33%。',
        '糖豆智能分拣原型：使用 Arduino UNO、TCS34725 与双舵机构建分拣装置；3 分钟分拣 76 颗，准确率 97.4%。',
        '机器人与自主导航：基于云深处四足机器人平台集成 ROS Navigation、Cartographer 与 RViz，完成实时建图、传感器可视化和多点导航演示。',
    ], 7.25)
    section_heading(doc, '工具、语言与专业能力')
    body_text(doc, '项目管理：里程碑、风险识别、跨部门协同、供应商管理、测试验证、试产/量产、客诉闭环', 7.25)
    body_text(doc, '智能硬件：门禁整机、门锁与电控锁、固件与协议、局域网/云平台基础、硬件失效分析、结构与装配', 7.25)
    body_text(doc, '工具：SolidWorks、MATLAB/Simulink、Python、Altium Designer、LabVIEW、C++、TensorFlow/PyTorch', 7.25)
    body_text(doc, '语言：CET-6 519｜CET-4 520｜IELTS 6.0｜可进行英文需求澄清、技术沟通与项目协同', 7.25)
    body_text(doc, '长期习惯：坚持每日跑步约 4 km；2025—2026 年跑步记录累计超过 1,600 km。', 7.25)
    dark_bar(doc, '查看完整产品实拍、项目视频与双语履历', 'mindscribe-app.pages.dev →', True)

    enlarge_document_typography(doc, 1.08)
    path = OUT / 'Liu-Moxin-Smart-Hardware-PM-CN.docx'
    doc.save(path)
    return path


def build_en() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.0)
    sec.bottom_margin = Cm(0.9)
    sec.left_margin = Cm(1.25)
    sec.right_margin = Cm(1.25)
    sec.header_distance = Cm(0)
    sec.footer_distance = Cm(0)
    styles = doc.styles
    styles['Normal'].font.name = 'Arial'
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    styles['Normal'].font.size = Pt(8)

    add_header(doc, name='MOXIN LIU', title='Smart Hardware Project Manager', subtitle=None, location='Hangzhou, China · +86 159 9446 1673', email='scretliasion@163.com', portfolio_label='Portfolio: mindscribe-app.pages.dev', language_line='English working proficiency')
    body_text(doc, '', after=1)
    add_profile_box(doc, 'PROFILE', 'Smart hardware project manager experienced in the full lifecycle of access control terminals, residential smart locks, glass-door locks, and RV locks—from requirements and R&D tracking through verification, compliance, supplier coordination, pilot builds, mass production, and after-sales closure. Completed an English-taught master’s program and use English as a working language, with independent client coordination across Nepal, Pakistan, and other overseas markets. Able to bridge hardware, firmware, mechanics, cloud/LAN architecture, quality, and supply chain teams to deliver products against milestones and risks.')
    section_heading(doc, 'CORE CAPABILITIES')
    add_caps(doc, 'Product Lifecycle  ·  Cross-functional Delivery  ·  HW/FW Integration  ·  Verification  ·  Compliance  ·  NPI & Mass Production  ·  Issue Closure  ·  English Client Communication')
    section_heading(doc, 'PROFESSIONAL EXPERIENCE')
    role_heading(doc, 'Deli Group', 'Smart Hardware Project Manager', 'Mar 2026 — Present')
    body_text(doc, 'Own project execution across access control terminals, residential smart locks, glass-door locks, and RV locks, covering in-house products and OEM programs.', 7.45, MUTED, after=1.0)
    bullet_list(doc, [
        'Managed concurrent 2.4-inch, 5-inch, and 7-inch access control projects and OEM variants; completed development, verification, and pilot builds for AL912C, AC213, and AC205C, supporting their August 2026 launch.',
        'Led type testing, hardware reliability verification, first-round functional validation, and packaging design for waterproof AC516/AL962C; organised retesting and risk reviews for palm-vein recognition and component lead times.',
        'Drove new-standard compliance for the AC22 face-recognition lock by comparing certification paths and coordinating corrective actions; aligned legacy products including DL-AC27 with updated requirements.',
        'Managed the RV lock from patent-risk screening and supplier sourcing to pricing, pilot verification, and on-site production follow-up; coordinated a low-temperature latch fix validated at -30°C.',
        'Took over a 10-item historical after-sales backlog and closed 5 cases; coordinated root-cause analysis and corrective action across soldering, firmware, recognition, structural, and power issues.',
        'Produced a failure-analysis brief for AC512 doorbell response issues by mapping circuit logic and common wiring errors, improving field diagnosis and after-sales consistency.',
        'Standardised product parameters, operating instructions, specifications, and operation-video QR codes; clarified ownership and handoff rules between product and after-sales teams.',
        'Independently coordinated with customers in Nepal, Pakistan, and other overseas markets in English across requirements, solutions, customisation, and project progress.',
    ], 7.25)
    dark_bar(doc, 'DELIVERY METHOD', 'Milestones → risk register → owners → acceptance criteria → verified closure')
    role_heading(doc, 'Guangxi Nuote Broadcasting Co., Ltd.', 'Thermal Fusion Equipment Project Intern', 'Jun 2024 — Sep 2024')
    bullet_list(doc, ['Coordinated design and factory resources across the product development cycle and helped move three new products from design into pilot production.'], 7.25)
    product_matrix(doc, [
        ('Access Terminals', '2.4 / 5 / 7-inch and OEM'),
        ('Residential Locks', 'Delivery and compliance'),
        ('Glass-door Locks', 'Recognition and interaction'),
        ('RV Locks', 'Mechanism, pilot, production'),
    ])
    dark_bar(doc, 'DELIVERY CHAIN', 'Definition → R&D alignment → verification → supply chain → production → issue closure')
    target_footer(doc, 'Target roles: Smart Hardware Project Manager · Product Project Manager · Access Control / IoT / Robotics & Automation')

    doc.add_page_break()
    add_page2_header(doc, 'MOXIN LIU', 'Smart Hardware Project Manager', '+86 159 9446 1673 · scretliasion@163.com')
    metric_row(doc, [('3 models', '2.4-inch terminals launched'), ('5 / 10', 'historical issues closed'), ('-30°C', 'RV lock validation passed')])
    section_heading(doc, 'SELECTED DELIVERY CASES')
    add_case(doc, '01 | Multi-model Access Control Launches & OEM Delivery', 'Structured project gates from planning and protocol/functional tests through reliability verification, pilot builds, and warehouse release. Coordinated R&D, quality, sourcing, suppliers, and production across in-house terminals and OEM firmware/documentation variants.', 'AL912C / AC213 / AC205C / AC516 / AL962C / AC611 / AC711 / AC612 / AC615')
    add_case(doc, '02 | Smart Lock Compliance Transition', 'Managed the GB 21556.2-2025 transition for new and in-market products, including AC22 certification and upgrades across AC10, AC21, AC310, and DL-AC27. Balanced launch timing with corrective action, documentation updates, and legacy compliance.')
    add_case(doc, '03 | Quality & Customer Issue Closure', 'Established a repeatable flow covering reproduction, impact assessment, hardware/firmware isolation, ownership, verification, and closure. Applied it to PCB soldering, firmware, recognition, structural, power, and installation wiring cases.')
    section_heading(doc, 'EDUCATION')
    role_heading(doc, 'Zhejiang Gongshang University Sussex Artificial Intelligence Institute', 'M.Sc. in Robotics and Autonomous Systems — English-taught', '')
    body_text(doc, 'Machine Learning & Deep Learning, Computer Vision, Motion Planning & Control, Embedded AI, 3D Modelling and Mechanical Simulation', 7.2, MUTED)
    role_heading(doc, 'Zhejiang Gongshang University', 'Bachelor’s Degree in Electronic Information', '')
    body_text(doc, 'Circuit Analysis, Digital/Analogue Electronics, Microcontrollers, Signals & Systems, EDA', 7.2, MUTED)
    section_heading(doc, 'ENGINEERING & RESEARCH PROJECTS')
    bullet_list(doc, [
        'High-speed railway 5G beam/resource coordination: four-scenario platform using PPO and Q-learning; report records straight-track handover success improving from 92.8% to 97.4%.',
        'UAV cellular blind-spot search: path planning, channel modelling, and signal visualisation; inward-spiral paths reduced simulated mission time by approximately 28.92%–33.33%.',
        'Candy colour-sorting prototype: Arduino UNO, TCS34725 sensor, and dual servos; 76 pieces sorted in three minutes at 97.4% accuracy.',
        'Robotics and autonomous navigation: integrated ROS Navigation, Cartographer, and RViz on a DEEP Robotics quadruped for live mapping, sensor visualisation, and multi-point navigation.',
    ], 7.1)
    section_heading(doc, 'TOOLS, LANGUAGES & DOMAIN KNOWLEDGE')
    body_text(doc, 'Project Delivery: milestones, risk registers, cross-functional coordination, suppliers, verification, pilot/mass production, issue closure', 7.1)
    body_text(doc, 'Smart Hardware: access control, smart locks, firmware/protocols, LAN/cloud fundamentals, failure analysis, mechanical integration', 7.1)
    body_text(doc, 'Tools: SolidWorks, MATLAB/Simulink, Python, Altium Designer, LabVIEW, C++, TensorFlow/PyTorch', 7.1)
    body_text(doc, 'Languages: Mandarin (native); English (working proficiency) — CET-6 519, CET-4 520, IELTS 6.0', 7.1)
    body_text(doc, 'Long-term Habit: run approximately 4 km daily; more than 1,600 km recorded across 2025–2026.', 7.1)
    dark_bar(doc, 'View product fieldwork, project videos, and the bilingual portfolio', 'mindscribe-app.pages.dev →', True)

    enlarge_document_typography(doc, 1.08)
    path = OUT / 'Liu-Moxin-Smart-Hardware-PM-EN.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    for output in (build_cn(), build_en()):
        print(output)
